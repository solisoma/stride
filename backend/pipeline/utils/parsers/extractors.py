from pipeline.utils.parsers.page_types import SourcePage
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path


def analyze_images(images, vision):
    if not images:
        return ''
    with ThreadPoolExecutor(max_workers=min(len(images), 5)) as executor:
        descriptions = list(executor.map(vision.extract_image_content, images))
    useful = [description for description in descriptions if description.strip()]
    if not useful:
        return ''
    return '\n\n' + '\n\n'.join(f'[Image-derived content; verify against the original image]\n{description}' for description in useful)


def extract_pdf_pages(file_path: str, vision, cache=None) -> list[SourcePage]:
    from pypdf import PdfReader
    pages: list[SourcePage] = []
    for number, page in enumerate(PdfReader(file_path).pages, 1):
        def read_page():
            return (page.extract_text() or '') + analyze_images([image.data for image in page.images], vision)
        text = cache.run('pdf_page', number, read_page) if cache else read_page()
        if text.strip():
            pages.append({'page': number, 'text': text.strip()})
    return pages


def extract_docx_pages(file_path: str, vision, cache=None) -> list[SourcePage]:
    from docx import Document
    import xml.etree.ElementTree as ET
    doc = Document(file_path)
    paragraphs = []
    for number, paragraph in enumerate(doc.paragraphs):
        def read_paragraph():
            blobs = []
            for run in paragraph.runs:
                root = ET.fromstring(run.element.xml)
                for blip in root.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}blip'):
                    embed = blip.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                    if embed:
                        blobs.append(paragraph.part.related_parts[embed].blob)
            return paragraph.text + analyze_images(blobs, vision)
        text = cache.run('docx_paragraph', number, read_paragraph) if cache else read_paragraph()
        if text.strip():
            paragraphs.append(text.strip())
    # Retain table text in addition to the original paragraph/image extraction.
    paragraphs.extend(' | '.join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)
    return [{'page': f'section {i // 15 + 1}', 'text': '\n'.join(paragraphs[i:i + 15])} for i in range(0, len(paragraphs), 15)]


def extract_text_pages(file_path: str) -> list[SourcePage]:
    from pipeline.utils.core.text_splitter import RecursiveCharacterTextSplitter
    text = Path(file_path).read_text(encoding='utf-8')
    chunks = RecursiveCharacterTextSplitter(chunk_size=6000, chunk_overlap=0).split_text(text)
    return [{'page': f'section {i + 1}', 'text': text} for i, text in enumerate(chunks) if text.strip()]
